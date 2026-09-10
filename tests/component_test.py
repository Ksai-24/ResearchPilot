"""Test file parsing + export endpoints + server wiring (no LLM needed)."""
import asyncio
import sys

from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


async def main():
    from io import BytesIO
    ok = True

    print("== 1. parsers: txt / docx / pdf ==")
    from aira.files import extract
    r = extract("notes.txt", "hello world\nthis is a test file".encode())
    assert r["ok"] and "test file" in r["text"], r
    print("  txt OK")

    from docx import Document
    d = Document()
    d.add_paragraph("Aira research note.")
    d.add_paragraph("Second paragraph about herd immunity.")
    bio = BytesIO(); d.save(bio)
    r = extract("note.docx", bio.getvalue())
    assert r["ok"] and "herd immunity" in r["text"], r
    print("  docx OK")

    from fpdf import FPDF
    p = FPDF(); p.add_page(); p.set_font("helvetica", size=12)
    p.cell(0, 10, "PDF test document content.", new_x="LMARGIN")
    r = extract("note.pdf", bytes(p.output()))
    assert r["ok"] and "PDF test" in r["text"], r
    print("  pdf OK")

    print("\n== 2. export: txt / docx / pdf ==")
    from aira.export import export_chat
    msgs = [
        {"role": "user", "text": "What is herd immunity? — arrows → and “quotes” test", "files": []},
        {"role": "assistant", "text": "Indirect protection [1].",
         "sources": ["WHO — https://who.int/news"], "tools": ["Searching the web", "Reading a source"]},
    ]
    for fmt, check in (("txt", b"herd immunity"), ("docx", b"word/document.xml"), ("pdf", b"%PDF")):
        data, fname, mime = export_chat("Test chat", "RESEARCH", msgs, fmt)
        assert check in data[:2000] or check == b"%PDF" and data[:5] == b"%PDF", f"{fmt} bad: {fname}"
        print(f"  {fmt} OK  ({fname}, {len(data)} bytes, {mime})")

    print("\n== 3. server module imports ==")
    import aira.server
    print("  aira.server OK — routes:",
          [r.path for r in aira.server.app.routes if hasattr(r, 'path') and r.path.startswith('/api')])

    print("\nALL COMPONENT TESTS PASSED")
    return 0 if ok else 1


sys.exit(asyncio.run(main()))
